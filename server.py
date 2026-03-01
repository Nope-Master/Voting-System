#!/usr/bin/env python3
"""
NextCR Voting System - Server
Run: python server.py
Open: http://localhost:8080

Data stored in database.xlsx
"""

import http.server
import socketserver
import json
import os
import sys
import hashlib
import urllib.parse
import traceback

try:
    import openpyxl
except ImportError:
    print("=" * 50)
    print("  openpyxl not installed!")
    print("  Run: pip install openpyxl")
    print("=" * 50)
    sys.exit(1)

PORT = int(os.environ.get("PORT", 8080))
HOST = os.environ.get("HOST", "0.0.0.0")
DB_FILE = os.environ.get("DB_FILE", "database.xlsx")
PHOTO_DIR = "static/images/uploads"

# Create directories
os.makedirs("static/images", exist_ok=True)
os.makedirs(PHOTO_DIR, exist_ok=True)


def init_database():
    """Create database.xlsx with all required sheets if it doesn't exist or is corrupted."""
    need_create = False

    if not os.path.exists(DB_FILE):
        need_create = True
    else:
        try:
            wb = openpyxl.load_workbook(DB_FILE)
            wb.close()
        except Exception as e:
            print(f"  [WARN] database.xlsx corrupted: {e}")
            print(f"  [WARN] Recreating database...")
            try:
                os.remove(DB_FILE)
            except Exception:
                pass
            need_create = True

    if need_create:
        wb = openpyxl.Workbook()
        # Remove default sheet
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]
    else:
        wb = openpyxl.load_workbook(DB_FILE)

    # Define all sheets and their headers
    sheets = {
        "ValidStudentIDs": ["student_id", "name"],
        "Students": [
            "id", "name", "student_id", "password_hash",
            "contact_no", "photo", "voted_round1", "voted_round2",
            "profile_details"
        ],
        "Admins": [
            "id", "name", "password_hash", "admin_pin_hash",
            "contact_no", "photo", "profile_details"
        ],
        "Candidates": [
            "id", "name", "student_id", "round1_votes",
            "round2_votes", "qualified_for_final", "bio"
        ],
        "VotingSettings": ["key", "value"],
        "CurrentCR": ["key", "value"],
        "AdminUnlockPin": ["pin_hash"],
    }

    for sheet_name, headers in sheets.items():
        if sheet_name not in wb.sheetnames:
            ws = wb.create_sheet(sheet_name)
            for col_idx, header in enumerate(headers, 1):
                ws.cell(row=1, column=col_idx, value=header)
            print(f"  [DB] Created sheet: {sheet_name}")

    # Insert default VotingSettings if empty
    ws = wb["VotingSettings"]
    if ws.max_row is None or ws.max_row < 2:
        defaults = [
            ("current_round", "0"),
            ("status", "NotStarted"),
            ("qualification_limit", "5"),
            ("semi_results_declared", "false"),
            ("final_results_declared", "false"),
            ("round_start_time", ""),
            ("round_duration", ""),
        ]
        for i, (k, v) in enumerate(defaults, 2):
            ws.cell(row=i, column=1, value=k)
            ws.cell(row=i, column=2, value=v)

    # Insert default CurrentCR if empty
    ws = wb["CurrentCR"]
    if ws.max_row is None or ws.max_row < 2:
        defaults = [
            ("name", "Not Assigned Yet"),
            ("photo", ""),
            ("contact_no", ""),
            ("details", ""),
            ("profile_details", "[]"),
        ]
        for i, (k, v) in enumerate(defaults, 2):
            ws.cell(row=i, column=1, value=k)
            ws.cell(row=i, column=2, value=v)

    try:
        wb.save(DB_FILE)
        print(f"  [DB] {DB_FILE} ready with {len(sheets)} sheets")
    except Exception as e:
        print(f"  [ERROR] Could not save {DB_FILE}: {e}")
    finally:
        wb.close()


def safe_str(val):
    """Convert cell value to string safely."""
    if val is None:
        return ""
    return str(val).strip()


def safe_int(val):
    """Convert cell value to int safely."""
    if val is None:
        return 0
    try:
        return int(float(str(val)))
    except (ValueError, TypeError):
        return 0


def safe_bool(val):
    """Convert cell value to boolean safely."""
    if val is None:
        return False
    return str(val).strip().lower() == "true"


def safe_json(val):
    """Parse JSON from cell value safely."""
    if val is None:
        return []
    try:
        result = json.loads(str(val))
        if isinstance(result, list):
            return result
        return []
    except (json.JSONDecodeError, TypeError):
        return []


def truncate_for_excel(val, max_len=32000):
    """Truncate string to fit in Excel cell. Photos that are too long get cleared."""
    if val is None:
        return ""
    s = str(val)
    if len(s) > max_len:
        # If it's a base64 photo, save to file instead
        if s.startswith("data:image"):
            return ""  # Will be handled by photo saving
        return s[:max_len]
    return s


def save_photo_to_file(base64_data, filename_prefix):
    """Save base64 photo data to a file and return the file path."""
    if not base64_data or not base64_data.startswith("data:image"):
        return base64_data  # Not a base64 image, return as-is

    try:
        # Extract the actual base64 content
        import base64 as b64
        header, data = base64_data.split(",", 1)

        # Determine extension
        ext = "jpg"
        if "png" in header:
            ext = "png"
        elif "gif" in header:
            ext = "gif"
        elif "webp" in header:
            ext = "webp"

        filename = f"{filename_prefix}.{ext}"
        filepath = os.path.join(PHOTO_DIR, filename)

        with open(filepath, "wb") as f:
            f.write(b64.b64decode(data))

        return f"/{filepath}"
    except Exception as e:
        print(f"  [WARN] Failed to save photo: {e}")
        return ""


def load_full_db():
    """Load entire database from Excel as a JSON-serializable dict."""
    wb = openpyxl.load_workbook(DB_FILE, read_only=True, data_only=True)

    try:
        # --- ValidStudentIDs ---
        valid_ids = []
        ws = wb["ValidStudentIDs"]
        for row in ws.iter_rows(min_row=2, max_col=2):
            sid = safe_str(row[0].value)
            if sid:
                name = safe_str(row[1].value) if len(row) > 1 else ""
                valid_ids.append({"id": sid, "name": name})

        # --- Students ---
        students = []
        ws = wb["Students"]
        for row in ws.iter_rows(min_row=2, max_col=9):
            uid = safe_str(row[0].value)
            if not uid:
                continue
            students.append({
                "id": uid,
                "name": safe_str(row[1].value),
                "student_id": safe_str(row[2].value),
                "password_hash": safe_str(row[3].value),
                "contact_no": safe_str(row[4].value),
                "photo": safe_str(row[5].value),
                "voted_round1": safe_bool(row[6].value),
                "voted_round2": safe_bool(row[7].value),
                "profile_details": safe_json(row[8].value),
            })

        # --- Admins ---
        admins = []
        ws = wb["Admins"]
        for row in ws.iter_rows(min_row=2, max_col=7):
            uid = safe_str(row[0].value)
            if not uid:
                continue
            admins.append({
                "id": uid,
                "name": safe_str(row[1].value),
                "password_hash": safe_str(row[2].value),
                "admin_pin_hash": safe_str(row[3].value),
                "contact_no": safe_str(row[4].value),
                "photo": safe_str(row[5].value),
                "profile_details": safe_json(row[6].value),
            })

        # --- Candidates ---
        candidates = []
        ws = wb["Candidates"]
        for row in ws.iter_rows(min_row=2, max_col=7):
            uid = safe_str(row[0].value)
            if not uid:
                continue
            candidates.append({
                "id": uid,
                "name": safe_str(row[1].value),
                "student_id": safe_str(row[2].value),
                "round1_votes": safe_int(row[3].value),
                "round2_votes": safe_int(row[4].value),
                "qualified_for_final": safe_bool(row[5].value),
                "bio": safe_str(row[6].value),
            })

        # --- VotingSettings (key-value) ---
        vs = {}
        ws = wb["VotingSettings"]
        for row in ws.iter_rows(min_row=2, max_col=2):
            k = safe_str(row[0].value)
            v = safe_str(row[1].value)
            if k:
                vs[k] = v

        rst = vs.get("round_start_time", "")
        rd = vs.get("round_duration", "")
        try:
            rst = int(float(rst)) if rst else None
        except (ValueError, TypeError):
            rst = None
        try:
            rd = int(float(rd)) if rd else None
        except (ValueError, TypeError):
            rd = None

        voting_settings = {
            "current_round": safe_int(vs.get("current_round", "0")),
            "status": vs.get("status", "NotStarted"),
            "qualification_limit": safe_int(vs.get("qualification_limit", "5")) or 5,
            "semi_results_declared": vs.get("semi_results_declared", "false") == "true",
            "final_results_declared": vs.get("final_results_declared", "false") == "true",
            "round_start_time": rst,
            "round_duration": rd,
        }

        # --- CurrentCR (key-value) ---
        cr = {}
        ws = wb["CurrentCR"]
        for row in ws.iter_rows(min_row=2, max_col=2):
            k = safe_str(row[0].value)
            v = safe_str(row[1].value)
            if k:
                cr[k] = v

        current_cr = {
            "name": cr.get("name", "Not Assigned Yet"),
            "photo": cr.get("photo", ""),
            "contact_no": cr.get("contact_no", ""),
            "details": cr.get("details", ""),
            "profile_details": safe_json(cr.get("profile_details", "[]")),
        }

        # --- AdminUnlockPin ---
        pin_hash = ""
        ws = wb["AdminUnlockPin"]
        for row in ws.iter_rows(min_row=2, max_col=1):
            val = safe_str(row[0].value)
            if val:
                pin_hash = val
                break

        return {
            "ValidStudentIDs": valid_ids,
            "Students": students,
            "Admins": admins,
            "Candidates": candidates,
            "VotingSettings": voting_settings,
            "CurrentCR": current_cr,
            "AdminUnlockPin": pin_hash,
        }

    finally:
        wb.close()


def save_full_db(db):
    """Save entire database dict back to Excel."""
    wb = openpyxl.load_workbook(DB_FILE)

    try:
        # --- Helper to clear data rows (keep header in row 1) ---
        def clear_data(ws):
            if ws.max_row and ws.max_row > 1:
                # delete_rows(start, count) - delete (max_row - 1) rows starting from row 2
                ws.delete_rows(2, ws.max_row - 1)

        # --- ValidStudentIDs ---
        ws = wb["ValidStudentIDs"]
        clear_data(ws)
        for i, v in enumerate(db.get("ValidStudentIDs", []), 2):
            ws.cell(row=i, column=1, value=safe_str(v.get("id")))
            ws.cell(row=i, column=2, value=safe_str(v.get("name")))

        # --- Students ---
        ws = wb["Students"]
        clear_data(ws)
        for i, s in enumerate(db.get("Students", []), 2):
            ws.cell(row=i, column=1, value=safe_str(s.get("id")))
            ws.cell(row=i, column=2, value=safe_str(s.get("name")))
            ws.cell(row=i, column=3, value=safe_str(s.get("student_id")))
            ws.cell(row=i, column=4, value=safe_str(s.get("password_hash")))
            ws.cell(row=i, column=5, value=safe_str(s.get("contact_no")))

            # Handle photo - save to file if base64
            photo = s.get("photo", "")
            if photo and str(photo).startswith("data:image"):
                photo = save_photo_to_file(photo, f"student_{s.get('id', 'unknown')}")
            ws.cell(row=i, column=6, value=safe_str(photo))

            ws.cell(row=i, column=7, value=str(bool(s.get("voted_round1", False))).lower())
            ws.cell(row=i, column=8, value=str(bool(s.get("voted_round2", False))).lower())
            ws.cell(row=i, column=9, value=json.dumps(s.get("profile_details", [])))

        # --- Admins ---
        ws = wb["Admins"]
        clear_data(ws)
        for i, a in enumerate(db.get("Admins", []), 2):
            ws.cell(row=i, column=1, value=safe_str(a.get("id")))
            ws.cell(row=i, column=2, value=safe_str(a.get("name")))
            ws.cell(row=i, column=3, value=safe_str(a.get("password_hash")))
            ws.cell(row=i, column=4, value=safe_str(a.get("admin_pin_hash")))
            ws.cell(row=i, column=5, value=safe_str(a.get("contact_no")))

            photo = a.get("photo", "")
            if photo and str(photo).startswith("data:image"):
                photo = save_photo_to_file(photo, f"admin_{a.get('id', 'unknown')}")
            ws.cell(row=i, column=6, value=safe_str(photo))

            ws.cell(row=i, column=7, value=json.dumps(a.get("profile_details", [])))

        # --- Candidates ---
        ws = wb["Candidates"]
        clear_data(ws)
        for i, c in enumerate(db.get("Candidates", []), 2):
            ws.cell(row=i, column=1, value=safe_str(c.get("id")))
            ws.cell(row=i, column=2, value=safe_str(c.get("name")))
            ws.cell(row=i, column=3, value=safe_str(c.get("student_id")))
            ws.cell(row=i, column=4, value=safe_int(c.get("round1_votes")))
            ws.cell(row=i, column=5, value=safe_int(c.get("round2_votes")))
            ws.cell(row=i, column=6, value=str(bool(c.get("qualified_for_final", False))).lower())
            ws.cell(row=i, column=7, value=safe_str(c.get("bio")))

        # --- VotingSettings ---
        ws = wb["VotingSettings"]
        clear_data(ws)
        vs = db.get("VotingSettings", {})
        kv = [
            ("current_round", str(vs.get("current_round", 0))),
            ("status", str(vs.get("status", "NotStarted"))),
            ("qualification_limit", str(vs.get("qualification_limit", 5))),
            ("semi_results_declared", str(bool(vs.get("semi_results_declared", False))).lower()),
            ("final_results_declared", str(bool(vs.get("final_results_declared", False))).lower()),
            ("round_start_time", str(vs.get("round_start_time") or "")),
            ("round_duration", str(vs.get("round_duration") or "")),
        ]
        for i, (k, v) in enumerate(kv, 2):
            ws.cell(row=i, column=1, value=k)
            ws.cell(row=i, column=2, value=v)

        # --- CurrentCR ---
        ws = wb["CurrentCR"]
        clear_data(ws)
        cr = db.get("CurrentCR", {})

        cr_photo = cr.get("photo", "")
        if cr_photo and str(cr_photo).startswith("data:image"):
            cr_photo = save_photo_to_file(cr_photo, "current_cr")

        cr_kv = [
            ("name", cr.get("name", "Not Assigned Yet")),
            ("photo", safe_str(cr_photo)),
            ("contact_no", cr.get("contact_no", "")),
            ("details", cr.get("details", "")),
            ("profile_details", json.dumps(cr.get("profile_details", []))),
        ]
        for i, (k, v) in enumerate(cr_kv, 2):
            ws.cell(row=i, column=1, value=k)
            ws.cell(row=i, column=2, value=v)

        # --- AdminUnlockPin ---
        ws = wb["AdminUnlockPin"]
        clear_data(ws)
        pin = db.get("AdminUnlockPin", "")
        if pin:
            ws.cell(row=2, column=1, value=str(pin))

        wb.save(DB_FILE)
        print(f"  [DB] Saved to {DB_FILE}")
        return True

    except Exception as e:
        print(f"  [ERROR] Save failed: {e}")
        traceback.print_exc()
        return False
    finally:
        wb.close()


class NextCRHandler(http.server.BaseHTTPRequestHandler):
    """HTTP Request Handler for NextCR."""

    # Increase max request size for photo uploads (10MB)
    MAX_CONTENT_LENGTH = 10 * 1024 * 1024

    def log_message(self, fmt, *args):
        """Custom log - only show API calls."""
        try:
            path = args[0].split()[1] if args else ""
        except (IndexError, AttributeError):
            path = ""
        if path.startswith("/api/"):
            code = args[1] if len(args) > 1 else ""
            print(f"  [API] {path} -> {code}")

    def send_json(self, data, code=200):
        """Send JSON response."""
        try:
            body = json.dumps(data, ensure_ascii=False).encode("utf-8")
        except Exception as e:
            body = json.dumps({"ok": False, "error": str(e)}).encode("utf-8")
            code = 500
        self.send_response(code)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Cache-Control", "no-cache, no-store")
        self.end_headers()
        self.wfile.write(body)

    def send_file(self, filepath, content_type):
        """Send a static file."""
        if not os.path.exists(filepath):
            self.send_response(404)
            self.send_header("Content-Type", "text/plain")
            self.end_headers()
            self.wfile.write(b"File not found")
            return
        try:
            with open(filepath, "rb") as f:
                data = f.read()
            self.send_response(200)
            self.send_header("Content-Type", content_type)
            self.send_header("Content-Length", str(len(data)))
            self.send_header("Cache-Control", "no-cache")
            self.end_headers()
            self.wfile.write(data)
        except Exception as e:
            self.send_response(500)
            self.send_header("Content-Type", "text/plain")
            self.end_headers()
            self.wfile.write(f"Error: {e}".encode())

    def read_body(self):
        """Read and parse JSON request body."""
        try:
            length = int(self.headers.get("Content-Length", 0))
            if length == 0:
                print("  [WARN] Empty request body")
                return {}
            if length > self.MAX_CONTENT_LENGTH:
                print(f"  [WARN] Request too large: {length} bytes")
                return {}
            body = self.rfile.read(length)
            text = body.decode("utf-8")
            result = json.loads(text)
            return result
        except json.JSONDecodeError as e:
            print(f"  [ERROR] JSON parse error: {e}")
            return {}
        except Exception as e:
            print(f"  [ERROR] read_body error: {e}")
            return {}

    def do_OPTIONS(self):
        """Handle CORS preflight."""
        self.send_response(200)
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")
        self.send_header("Access-Control-Max-Age", "86400")
        self.end_headers()

    def do_GET(self):
        """Handle GET requests."""
        path = urllib.parse.urlparse(self.path).path

        # Serve index.html
        if path == "/" or path == "/index.html":
            self.send_file("index.html", "text/html; charset=utf-8")
            return

        # Serve static files
        if path.startswith("/static/"):
            filepath = path.lstrip("/")
            # Security: prevent path traversal
            if ".." in filepath:
                self.send_response(403)
                self.end_headers()
                return
            ext_map = {
                ".css": "text/css; charset=utf-8",
                ".js": "application/javascript",
                ".png": "image/png",
                ".jpg": "image/jpeg",
                ".jpeg": "image/jpeg",
                ".gif": "image/gif",
                ".svg": "image/svg+xml",
                ".webp": "image/webp",
                ".ico": "image/x-icon",
                ".woff": "font/woff",
                ".woff2": "font/woff2",
            }
            ext = os.path.splitext(filepath)[1].lower()
            ct = ext_map.get(ext, "application/octet-stream")
            self.send_file(filepath, ct)
            return

        # API: Load database
        if path == "/api/db/load":
            try:
                db = load_full_db()
                self.send_json({"ok": True, "db": db})
            except Exception as e:
                print(f"  [ERROR] Load failed:")
                traceback.print_exc()
                self.send_json({"ok": False, "error": str(e)}, 500)
            return

        # 404
        self.send_response(404)
        self.send_header("Content-Type", "text/plain")
        self.end_headers()
        self.wfile.write(b"Not found")

    def do_POST(self):
        """Handle POST requests."""
        path = urllib.parse.urlparse(self.path).path

        # API: Save database
        if path == "/api/db/save":
            try:
                data = self.read_body()
                if not data:
                    self.send_json({"ok": False, "error": "Empty or invalid request body"}, 400)
                    return

                db = data.get("db")
                if not db or not isinstance(db, dict):
                    self.send_json({"ok": False, "error": "Missing 'db' field in request"}, 400)
                    return

                # Verify required keys exist
                required = ["ValidStudentIDs", "Students", "Admins", "Candidates",
                           "VotingSettings", "CurrentCR"]
                for key in required:
                    if key not in db:
                        self.send_json({"ok": False, "error": f"Missing key: {key}"}, 400)
                        return

                success = save_full_db(db)
                if success:
                    self.send_json({"ok": True})
                else:
                    self.send_json({"ok": False, "error": "Failed to write to Excel"}, 500)

            except Exception as e:
                print(f"  [ERROR] Save failed:")
                traceback.print_exc()
                self.send_json({"ok": False, "error": str(e)}, 500)
            return

        # 404
        self.send_response(404)
        self.send_header("Content-Type", "text/plain")
        self.end_headers()
        self.wfile.write(b"Not found")


class ReusableTCPServer(socketserver.TCPServer):
    """TCP Server that allows port reuse."""
    allow_reuse_address = True


if __name__ == "__main__":
    # Initialize database
    init_database()

    print()
    print("=" * 50)
    print("  NextCR Voting System")
    print(f"  http://{HOST}:{PORT}")
    print(f"  Database: {DB_FILE}")
    print("=" * 50)
    print()
    print("  Press Ctrl+C to stop")
    print()

    try:
        with ReusableTCPServer((HOST, PORT), NextCRHandler) as httpd:
            try:
                httpd.serve_forever()
            except KeyboardInterrupt:
                print("\n  Shutting down...")
                httpd.shutdown()
    except OSError as e:
        if "Address already in use" in str(e):
            print(f"\n  [ERROR] Port {PORT} is already in use!")
            print(f"  Try: kill the other process or change PORT in server.py")
        else:
            raise
